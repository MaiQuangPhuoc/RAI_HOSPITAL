"""PDF -> KDパンチ入力シート (demo).   Chạy:  python -m streamlit run app.py

Cấu trúc:  code/app.py , code/dic_output.txt , data/sheet.xlsx (chỉ có tới dòng 3)
Tự tạo thêm thư mục code/kd_sheet_component/ (component hiển thị + nhập liệu sheet).
"""
import ast
import colorsys
import io
import json
import re
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from functools import lru_cache
from html import escape
from pathlib import Path

import streamlit as st
import streamlit.components.v1 as components
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# ======================= CẤU HÌNH =======================
BASE = Path(__file__).resolve().parent                  # D:\rikai\RAI-LLM\code
DIC_PATH = BASE / 'dic_output.txt'
TEMPLATE_PATH = BASE.parent / 'data' / 'sheet.xlsx'     # D:\rikai\RAI-LLM\data\sheet.xlsx
SHEET_NAME = '入力シート'

SECONDS_PER_FILE = 3
FIRST_DATA_ROW = 4
FREEZE_ROWS = 3
FREEZE_COLS = 6
LEFT_PERCENT = 30
MAX_ROWS = 30                                           # UI chỉ hiển thị 30 dòng
RESULT_NAME = 'KDパンチ入力シート_ketqua.xlsx'

# Số trong tên file -> tên dict:  1 -> dic_4, 2 -> dic_6, 3 -> dic_8, 4 -> dic_10, 5 -> dic_12
DIC_BY_NUMBER = {n: f'dic_{2 + 2 * n}' for n in range(1, 6)}
# Dòng lẻ (5, 7, 9...): giá trị mặc định, người dùng click vào ô để sửa
ODD_DEFAULTS = {'回目': 2, '入力日': '4/9', '入力者': 'iida', 'Verify': 0}

GREEN, GREEN_DARK = '#217346', '#185c37'
# ========================================================


def find_template():
    if TEMPLATE_PATH.exists():
        return TEMPLATE_PATH
    raise FileNotFoundError(f'テンプレートExcelファイルが見つかりません: {TEMPLATE_PATH}')


def find_dic():
    if DIC_PATH.exists():
        return DIC_PATH
    raise FileNotFoundError(f'辞書ファイルが見つかりません: {DIC_PATH}')


def load_dicts(path):
    """Đọc dic_output.txt (dạng dic_4 = {...}) -> {'dic_4': {...}, ...}"""
    tree = ast.parse(Path(path).read_text(encoding='utf-8'))
    return {n.targets[0].id: ast.literal_eval(n.value)
            for n in tree.body if isinstance(n, ast.Assign)}


def file_number(filename):
    """abx_1.pdf -> 1 ; data-1.pdf -> 1 ; không có số -> None"""
    stem = Path(filename).stem
    found = re.findall(r'[_\-\s](\d+)', stem)
    if found:
        return int(found[-1])
    found = re.findall(r'\d+', stem)
    return int(found[0]) if found else None


def header_keys(ws):
    """Tên cột từ dòng 2/3, tên trùng thì thêm _2, _3."""
    row2 = [c.value for c in ws[2]]
    row3 = [c.value for c in ws[3]]
    zip_idx = next(i for i, v in enumerate(row2) if v == '郵便番号')
    names = [(row2[i] if row2[i] is not None else row3[i]) if i <= zip_idx else row2[i]
             for i in range(len(row2))]
    count, keys = {}, []
    for n in names:
        count[n] = count.get(n, 0) + 1
        keys.append(n if count[n] == 1 else f'{n}_{count[n]}')
    return keys


def is_editable_row(r):
    """Dòng lẻ trong vùng dữ liệu: 5, 7, 9..."""
    return r >= FIRST_DATA_ROW and (r - FIRST_DATA_ROW) % 2 == 1


# ---------------------------- Màu / viền từ file xlsx ----------------------------
def _lum(hexrgb):
    r, g, b = (int(hexrgb[i:i + 2], 16) for i in (0, 2, 4))
    return (0.299 * r + 0.587 * g + 0.114 * b) / 255


@lru_cache(maxsize=None)
def _tint(hexrgb, tint):
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
    return '%02x%02x%02x' % (round(r * 255), round(g * 255), round(b * 255))


def theme_palette(wb):
    xml = getattr(wb, 'loaded_theme', None)
    if not xml:
        return []
    s = xml.decode('utf-8', 'ignore') if isinstance(xml, bytes) else xml
    m = re.search(r'<a:clrScheme.*?</a:clrScheme>', s, re.S)
    if not m:
        return []
    d = dict(re.findall(r'<a:(dk1|lt1|dk2|lt2|accent[1-6])>\s*<a:(?:srgbClr val="|sysClr val="\w+" lastClr=")'
                        r'([0-9A-Fa-f]{6})', m.group(0)))
    order = ['lt1', 'dk1', 'lt2', 'dk2'] + [f'accent{i}' for i in range(1, 7)]
    return [d.get(k) for k in order]


def color_hex(color, pal):
    if color is None:
        return None
    if color.type == 'rgb' and isinstance(color.rgb, str):
        return None if color.rgb == '00000000' else color.rgb[-6:].lower()
    if color.type == 'theme' and pal and color.theme is not None and color.theme < len(pal) and pal[color.theme]:
        base = pal[color.theme].lower()
        return _tint(base, color.tint) if color.tint else base
    return None


BORDER_W = {'thin': '1px solid', 'hair': '1px solid', 'dotted': '1px dotted', 'dashed': '1px dashed',
            'dashDot': '1px dashed', 'dashDotDot': '1px dashed', 'medium': '2px solid',
            'mediumDashed': '2px dashed', 'mediumDashDot': '2px dashed', 'thick': '3px solid',
            'double': '3px double', 'slantDashDot': '2px dashed', 'mediumDashDotDot': '2px dashed'}


def border_css(side, b, pal):
    if b is None or not b.style:
        return None
    col = color_hex(b.color, pal)
    col = f'#{col}' if col and _lum(col) > 0.15 else 'var(--bord)'
    return f'border-{side}:{BORDER_W.get(b.style, "1px solid")} {col}'


# ---------------------------- Giao diện sheet kiểu Excel (2 khung) ----------------------------
THEMES = {
    'light': dict(bg='#ffffff', fg='#0a0a0a', grid='#c3c9cf', hdr='#eef1f3', hdr_fg='#3c4043',
                  band='#00ccff', border='#b8bec4', alt='#ffeeb8', bord='#000000', scroll='#a9b0b6'),
    'dark':  dict(bg='#1b1d1c', fg='#ffffff', grid='#4a4f4c', hdr='#2a2d2c', hdr_fg='#d0d4d2',
                  band='#0b7a99', border='#4a4f4c', alt='#54491f', bord='#d8d8d8', scroll='#6a706c'),
}
HDR_H = 32
RN_W = 54
MIN_ROW = 36


def excel_css(dark):
    t = THEMES['dark' if dark else 'light']
    return f"""<style>
:root{{--bg:{t['bg']};--fg:{t['fg']};--grid:{t['grid']};--hdr:{t['hdr']};--hdr-fg:{t['hdr_fg']};
      --band:{t['band']};--border:{t['border']};--alt:{t['alt']};--bord:{t['bord']};
      --green:{GREEN};--green-d:{GREEN_DARK};--frz:#8a9096}}
*{{box-sizing:border-box}}
html,body{{height:100%}}
body{{margin:0;background:var(--bg);color:var(--fg);display:flex;flex-direction:column;
     font:600 15px 'Segoe UI','Yu Gothic UI','Meiryo',Calibri,Arial,sans-serif}}
.fxbar{{height:46px;flex:none;display:flex;align-items:center;gap:12px;padding:0 14px;background:var(--hdr);
       border:1px solid var(--border);border-bottom:none;border-radius:12px 12px 0 0;font-size:16px;font-weight:700}}
.nb{{min-width:130px;height:32px;display:flex;align-items:center;justify-content:space-between;gap:8px;
    padding:0 12px;background:var(--bg);border:1px solid var(--border);border-radius:8px}}
.nb i{{font-style:normal;color:var(--hdr-fg);font-size:10px}}
.fxi{{color:var(--green);font:italic 700 18px serif}}
#fx{{flex:1;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}}

.panes{{flex:1;min-height:0;display:flex;border:1px solid var(--border);background:var(--bg)}}
.pane{{overflow-x:scroll;position:relative;background:var(--bg)}}
#pl{{width:{LEFT_PERCENT}%;flex:none;overflow-y:hidden}}
#pr{{flex:1;min-width:0;overflow-y:scroll}}
.split{{flex:none;width:12px;cursor:col-resize;background:var(--hdr);border-left:1px solid var(--border);
       border-right:1px solid var(--border);display:flex;align-items:center;justify-content:center;
       user-select:none;transition:background .15s}}
.split::after{{content:'';width:4px;height:56px;border-radius:2px;background:var(--hdr-fg);opacity:.55}}
.split:hover,.split.on{{background:var(--green)}}
.split:hover::after,.split.on::after{{background:#fff;opacity:1}}

table{{border-collapse:separate;border-spacing:0;table-layout:fixed}}
td,th{{border-right:1px solid var(--grid);border-bottom:1px solid var(--grid);padding:2px 8px;overflow:hidden;
      background:var(--bg);color:var(--fg);vertical-align:bottom;font-weight:600}}
td.hd{{background:var(--band)}}
td.alt{{background:var(--alt);border-bottom:2px solid var(--bord)}}
td.sel{{outline:3px solid var(--green);outline-offset:-3px}}
td.ed{{cursor:text}}
td.ed:hover{{outline:2px solid var(--green);outline-offset:-2px}}
td.ed[contenteditable="true"]{{outline:3px solid var(--green);outline-offset:-3px;background:var(--bg);
                              color:var(--fg);caret-color:var(--green)}}
td.chg{{font-style:italic;text-decoration:underline dotted var(--green)}}
td.fr{{position:sticky;z-index:3;font-weight:700}}
td.fb{{box-shadow:0 3px 0 var(--frz)}}
th{{background:var(--hdr);color:var(--hdr-fg);font-size:14px;font-weight:700;text-align:center}}
th.fr{{position:sticky;z-index:5}}
th.rn{{position:sticky;left:0;z-index:4;width:{RN_W}px}}
th.rn.fr{{z-index:6}}
th.fb{{box-shadow:0 3px 0 var(--frz)}}

.tab{{height:46px;flex:none;background:var(--hdr);border:1px solid var(--border);border-top:none;
     display:flex;align-items:center;gap:8px;padding-left:14px;border-radius:0 0 12px 12px}}
.tab .ic{{width:30px;text-align:center;color:var(--hdr-fg);font-size:18px}}
.tab .t{{background:var(--bg);color:var(--green);font-weight:700;font-size:15px;padding:11px 22px;
        border-bottom:4px solid var(--green);border-radius:8px 8px 0 0}}
::-webkit-scrollbar{{width:14px;height:14px}}
::-webkit-scrollbar-thumb{{background:{t['scroll']};border-radius:8px;border:3px solid var(--hdr)}}
::-webkit-scrollbar-thumb:hover{{background:var(--green)}}
::-webkit-scrollbar-track{{background:var(--hdr)}}
</style>"""


def sheet_to_html(ws, last_row):
    """Vẽ sheet (30 dòng) thành 2 khung. Dòng lẻ (5, 7, 9...) có thể click để nhập."""
    pal = theme_palette(ws.parent)
    n_rows, n_cols = MAX_ROWS, ws.max_column

    widths = {}
    for cd in ws.column_dimensions.values():
        for c in range(cd.min or 1, min(cd.max or cd.min or 1, n_cols) + 1):
            widths[c] = cd.width
    px = {c: int((widths.get(c) or 8.43) * 9 + 8) for c in range(1, n_cols + 1)}
    row_h = {}
    for r in range(1, n_rows + 1):
        h = ws.row_dimensions[r].height
        row_h[r] = max(MIN_ROW, round(h * 1.333)) if h else MIN_ROW

    def render(c_from, c_to, with_rn, tid):
        cols_range = range(c_from, c_to + 1)
        span, skip, virt = {}, set(), {}
        for m in ws.merged_cells.ranges:
            if m.max_col < c_from or m.min_col > c_to or m.min_row > n_rows:
                continue
            a, e = max(m.min_col, c_from), min(m.max_col, c_to)
            span[(m.min_row, a)] = (m.max_row - m.min_row + 1, e - a + 1)
            if m.min_col < c_from:
                virt[(m.min_row, a)] = (m.min_row, m.min_col)
            skip |= {(rr, cc) for rr in range(m.min_row, m.max_row + 1) for cc in range(a, e + 1)} - {(m.min_row, a)}

        total = (RN_W if with_rn else 0) + sum(px[c] for c in cols_range)
        cols = ('<col style="width:%dpx">' % RN_W if with_rn else '') + ''.join(
            f'<col style="width:{px[c]}px">' for c in cols_range)
        rows = []
        ths = ('<th class="rn fr corner"></th>' if with_rn else '') + ''.join(
            f'<th class="fr">{get_column_letter(c)}</th>' for c in cols_range)
        rows.append(f'<tr data-h="{HDR_H}" style="height:{HDR_H}px">{ths}</tr>')

        for r in range(1, n_rows + 1):
            fbot = ' fb' if r == FREEZE_ROWS else ''
            frz = ' fr' if r <= FREEZE_ROWS else ''
            tds = [f'<th class="rn{frz}{fbot}">{r}</th>'] if with_rn else []
            is_alt = is_editable_row(r) and r <= last_row
            for c in cols_range:
                if (r, c) in skip:
                    continue
                rs, cs = span.get((r, c), (1, 1))
                r2, c2 = r + rs - 1, c + cs - 1
                is_virtual = (r, c) in virt
                cell = ws.cell(*virt.get((r, c), (r, c)))
                al, ft, fl = cell.alignment, cell.font, cell.fill
                v = None if is_virtual else cell.value
                css, classes = [], []

                fill = None
                if fl and fl.fill_type == 'solid':
                    fill = color_hex(fl.fgColor, pal)
                    if fill is None and fl.fgColor.type == 'theme' and r <= 3:
                        classes.append('hd')
                if fill:
                    css.append(f'background:#{fill}')
                    css.append('color:' + ('#fff' if _lum(fill) < 0.45 else '#000'))
                fc = color_hex(ft.color, pal) if ft and ft.color is not None else None
                if fc and (fill or _lum(fc) > 0.3):
                    css.append(f'color:#{fc}')
                if ft and ft.bold:
                    css.append('font-weight:800')
                if is_alt:
                    classes.append('alt')
                if r2 <= FREEZE_ROWS:
                    classes.append('fr')
                    if r2 == FREEZE_ROWS:
                        classes.append('fb')

                pos = ws.cell(row=r, column=c)
                right = ws.cell(row=r, column=c2).border.right
                bottom = ws.cell(row=r2, column=c).border.bottom
                left = pos.border.left if (c == c_from or not ws.cell(row=r, column=c - 1).border.right.style) else None
                top = pos.border.top if (r == 1 or not ws.cell(row=r - 1, column=c).border.bottom.style) else None
                for side, bd in (('right', right), ('bottom', bottom), ('left', left), ('top', top)):
                    s = border_css(side, bd, pal)
                    if s:
                        css.append(s)

                ha = al.horizontal if al.horizontal not in (None, 'general') else \
                    ('right' if isinstance(v, (int, float, datetime)) else 'left')
                css.append(f'text-align:{ha}')
                css.append('vertical-align:' + {'center': 'middle', 'top': 'top'}.get(al.vertical, 'bottom'))
                css.append('white-space:pre-wrap' if al.wrap_text else 'white-space:nowrap')
                if isinstance(v, datetime):
                    v = f'{v.month}/{v.day}'
                raw = '' if v is None else str(v)
                text = escape(raw).replace('\n', '<br>')
                src_r, src_c = virt.get((r, c), (r, c))
                ref = f'{get_column_letter(src_c)}{src_r}'
                if is_virtual:
                    for m in ws.merged_cells.ranges:
                        if (m.min_row, m.min_col) == (src_r, src_c):
                            ref = m.coord
                elif rs > 1 or cs > 1:
                    ref = f'{get_column_letter(c)}{r}:{get_column_letter(c2)}{r2}'
                attr = (f' rowspan="{rs}"' if rs > 1 else '') + (f' colspan="{cs}"' if cs > 1 else '')
                if is_alt and not is_virtual and rs == 1 and cs == 1:      # ô nhập được
                    classes.append('ed')
                    attr += f' data-r="{r}" data-c="{c}" data-o="{escape(raw, quote=True)}"'
                cls = f' class="{" ".join(classes)}"' if classes else ''
                tds.append(f'<td{cls}{attr} data-ref="{ref}" style="{";".join(css)}">{text}</td>')
            rows.append(f'<tr data-h="{row_h[r]}" style="height:{row_h[r]}px">{"".join(tds)}</tr>')
        return (f'<table id="{tid}" style="width:{total}px"><colgroup>{cols}</colgroup>'
                f'<tbody>{"".join(rows)}</tbody></table>')

    left_to = min(FREEZE_COLS, n_cols)
    left_html = render(1, left_to, True, 'tl')
    right_html = render(left_to + 1, n_cols, False, 'tr') if n_cols > left_to else '<table id="tr"></table>'

    return ('<div class="fxbar"><div class="nb"><span id="nb">A1</span><i>▼</i></div>'
            '<div class="fxi">fx</div><div id="fx"></div></div>'
            f'<div class="panes" id="panes"><div class="pane" id="pl">{left_html}</div>'
            '<div class="split" id="sp" title="左右の比率をドラッグして調整"></div>'
            f'<div class="pane" id="pr">{right_html}</div></div>'
            f'<div class="tab"><span class="ic">＋</span><span class="ic">≡</span>'
            f'<div class="t">{escape(ws.title)}</div></div>')


# ---------------------------- Component: hiển thị sheet + nhập liệu (2 chiều) ----------------------------
COMPONENT_HTML = r"""<!doctype html><html><head><meta charset="utf-8"></head>
<body style="margin:0"><div id="root" style="display:contents"></div>
<script>
const P=(type,extra)=>window.parent.postMessage(Object.assign({isStreamlitMessage:true,type:type},extra||{}),'*');
const $=id=>document.getElementById(id);
let sig=null,edits={},FR=3,lastH=0;
const send=()=>P('streamlit:setComponentValue',{value:edits,dataType:'json'});
function setTheme(css){
  let style=$('theme-style');
  if(!style){style=document.createElement('style');style.id='theme-style';document.head.appendChild(style);}
  const text=css.replace(/^<style>/,'').replace(/<\/style>$/,'');
  if(style.textContent!==text){style.textContent=text;relayout();}
}

function layout(){                                   // ghi -> đọc 1 lượt -> ghi (không reflow lặp)
  const L=[...document.querySelectorAll('#tl tr')],R=[...document.querySelectorAll('#tr tr')];
  L.concat(R).forEach(t=>t.style.height=t.dataset.h+'px');
  const hs=L.map((t,i)=>Math.max(t.offsetHeight,R[i]?R[i].offsetHeight:0));
  L.forEach((t,i)=>{t.style.height=hs[i]+'px';if(R[i])R[i].style.height=hs[i]+'px';});
  let top=0;
  for(let i=0;i<=FR;i++){
    [L[i],R[i]].forEach(tr=>{if(tr)tr.querySelectorAll('.fr').forEach(x=>x.style.top=top+'px');});
    top+=hs[i]||0;
  }
}
let raf=0;const relayout=()=>{cancelAnimationFrame(raf);raf=requestAnimationFrame(layout);};
function fit(){
  let h=760;try{h=Math.max(560,Math.min(window.parent.innerHeight-130,900));}catch(e){}
  if(h!==lastH){lastH=h;P('streamlit:setFrameHeight',{height:h});}
}
function pick(td){
  document.querySelectorAll('td.sel').forEach(x=>x.classList.remove('sel'));
  td.classList.add('sel');$('nb').textContent=td.dataset.ref;$('fx').textContent=td.innerText;
}

document.addEventListener('scroll',e=>{                     // 2 khung cuộn dọc cùng nhau
  if(e.target.id==='pr'){const pl=$('pl');if(pl)pl.scrollTop=e.target.scrollTop;}},true);
document.addEventListener('wheel',e=>{
  const pl=e.target.closest&&e.target.closest('#pl');if(!pl)return;
  if(e.shiftKey)pl.scrollLeft+=e.deltaY;else $('pr').scrollTop+=e.deltaY;
  e.preventDefault();},{passive:false});

let drag=false;                                             // kéo thanh giữa
document.addEventListener('mousedown',e=>{
  if(e.target.id!=='sp')return;drag=true;e.target.classList.add('on');
  document.body.style.userSelect='none';e.preventDefault();});
document.addEventListener('mousemove',e=>{
  if(!drag)return;const r=$('panes').getBoundingClientRect();
  $('pl').style.width=Math.min(Math.max(e.clientX-r.left,180),r.width-260)+'px';});
document.addEventListener('mouseup',()=>{
  if(!drag)return;drag=false;$('sp').classList.remove('on');document.body.style.userSelect='';layout();});

document.addEventListener('click',e=>{                      // click ô: chọn; ô dòng lẻ: bắt đầu nhập
  const td=e.target.closest&&e.target.closest('#panes td');if(!td)return;pick(td);
  if(td.classList.contains('ed')&&td.contentEditable!=='true'){
    td.dataset.p=td.innerText;td.contentEditable='true';td.focus();
    const rg=document.createRange();rg.selectNodeContents(td);rg.collapse(false);
    const s=getSelection();s.removeAllRanges();s.addRange(rg);
  }});
document.addEventListener('input',e=>{
  if(e.target.classList&&e.target.classList.contains('ed'))$('fx').textContent=e.target.innerText;});
document.addEventListener('keydown',e=>{
  const td=e.target;if(!td.classList||!td.classList.contains('ed'))return;
  if(e.key==='Enter'){e.preventDefault();td.blur();}
  else if(e.key==='Escape'){td.innerText=td.dataset.p||'';td.blur();}});
document.addEventListener('focusout',e=>{                   // rời ô: lưu thay đổi, gửi về Python
  const td=e.target;if(!td.classList||!td.classList.contains('ed'))return;
  td.contentEditable='false';
  const t=td.innerText.replace(/\n+$/,''),k=td.dataset.r+','+td.dataset.c,before=edits[k];
  if(t!==(td.dataset.o||'')){edits[k]=t;td.classList.add('chg');}
  else{delete edits[k];td.classList.remove('chg');}
  if(edits[k]!==before)send();});

window.addEventListener('message',e=>{
  if(e.data==='kd-scroll-top'){                             // nút ↑ ở góc phải dưới
    const pr=$('pr'),pl=$('pl');if(pr)pr.scrollTo({top:0,behavior:'smooth'});
    if(pl)pl.scrollTo({top:0,behavior:'smooth'});return;}
  const d=e.data;if(!d||d.type!=='streamlit:render')return;
  const a=d.args;
  setTheme(a.css);                                          // đổi theme mà không dựng lại/mất edits
  if(a.sig!==sig){                                          // chỉ dựng lại khi đổi kết quả
    sig=a.sig;FR=a.fr;edits=Object.assign({},a.init||{});
    $('root').innerHTML=a.html;
    for(const k in edits){const p=k.split(',');
      const td=document.querySelector('td.ed[data-r="'+p[0]+'"][data-c="'+p[1]+'"]');
      if(td){td.innerText=edits[k];td.classList.add('chg');}}
    const first=document.querySelector('td[data-ref]');if(first)pick(first);
    fit();layout();
    if(document.fonts)document.fonts.ready.then(relayout);
  }else fit();
});
window.addEventListener('resize',()=>{fit();relayout();});
P('streamlit:componentReady',{apiVersion:1});
</script></body></html>"""

COMP_DIR = BASE / 'kd_sheet_component'
COMP_DIR.mkdir(exist_ok=True)
_idx = COMP_DIR / 'index.html'
if not _idx.exists() or _idx.read_text(encoding='utf-8') != COMPONENT_HTML:
    _idx.write_text(COMPONENT_HTML, encoding='utf-8')
kd_sheet = components.declare_component('kd_sheet', path=str(COMP_DIR))


# ---------------------------- Xem nội dung PDF + nút lên đầu trang ----------------------------
def pdf_media_url(uploaded_file, container):
    """Đăng ký PDF qua media endpoint để link_button mở viewer của trình duyệt."""
    from streamlit import runtime

    return runtime.get_instance().media_file_mgr.add(
        uploaded_file.getvalue(),
        'application/pdf',
        container._get_delta_path_str(),
    )


TOP_BTN = """<script>
(function(){
  const d=window.parent.document;
  if(d.getElementById('kd-top-btn'))return;
  const b=d.createElement('button');
  b.id='kd-top-btn';b.title='先頭へ';b.textContent='↑';
  b.style.cssText='position:fixed;right:28px;bottom:28px;width:54px;height:54px;border-radius:50%;border:none;'
    +'background:#217346;color:#fff;font-size:26px;font-weight:700;cursor:pointer;z-index:999999;'
    +'box-shadow:0 6px 16px rgba(0,0,0,.35);transition:transform .15s,background .15s';
  b.onmouseenter=()=>{b.style.background='#185c37';b.style.transform='translateY(-3px)';};
  b.onmouseleave=()=>{b.style.background='#217346';b.style.transform='none';};
  b.onclick=()=>{
    const m=d.querySelector('[data-testid="stMain"]')||d.querySelector('section.main');
    if(m)m.scrollTo({top:0,behavior:'smooth'});
    d.querySelectorAll('iframe').forEach(f=>{try{f.contentWindow.postMessage('kd-scroll-top','*');}catch(e){}});
  };
  d.body.appendChild(b);
})();
</script>"""


# ---------------------------- Giao diện Streamlit ----------------------------
def app_css(dark):
    if dark:
        bg, fg, card, border, muted = '#121715', '#f2f6f3', '#1c2420', '#2f3b34', '#a9bbb0'
    else:
        bg, fg, card, border, muted = '#f4f7f5', '#13211a', '#ffffff', '#d5e0d9', '#5d6f64'
    return f"""<style>
.stApp{{background:{bg}}}
.stApp, .stApp p, .stApp label, .stApp h1, .stApp h2, .stApp h3, .stApp span, .stApp small{{color:{fg}}}
header[data-testid="stHeader"]{{background:transparent}}
.block-container{{padding:1.5rem 2rem 3rem;max-width:100%}}
h1{{font-weight:700;letter-spacing:-0.5px}}

[data-testid="stFileUploaderDropzone"]{{background:{card};border:2px dashed {GREEN};border-radius:14px;padding:1.2rem}}
[data-testid="stFileUploaderDropzone"] small{{color:{muted}}}
[data-testid="stFileUploaderFile"]{{background:{card};border:1px solid {border};border-radius:10px;padding:.2rem .6rem}}

.stButton > button, .stDownloadButton > button,
[data-testid="stFileUploaderDropzone"] button{{
  background:{GREEN};color:#fff !important;border:none;border-radius:10px;
  font-weight:600;padding:.55rem 1.3rem;box-shadow:0 2px 6px rgba(33,115,70,.25);transition:all .15s}}
.stButton > button *, .stDownloadButton > button *,
[data-testid="stFileUploaderDropzone"] button *{{color:#fff !important}}
.stButton > button:hover:not(:disabled), .stDownloadButton > button:hover:not(:disabled),
[data-testid="stFileUploaderDropzone"] button:hover{{background:{GREEN_DARK};transform:translateY(-1px);
  box-shadow:0 4px 10px rgba(33,115,70,.35)}}
.stButton > button:disabled{{background:{border};color:{muted} !important;box-shadow:none;cursor:not-allowed}}
.stButton > button:disabled *{{color:{muted} !important}}

[class*="st-key-view_"] a{{
  display:flex;background:{card} !important;border:2px solid {GREEN} !important;
  box-shadow:none !important;color:{GREEN} !important;transition:background .15s,color .15s}}
[class*="st-key-view_"] a *,
[class*="st-key-view_"] a *:visited{{color:{GREEN} !important}}
[class*="st-key-view_"] a:hover,
[class*="st-key-view_"] a:focus,
[class*="st-key-view_"] a:focus-visible,
[class*="st-key-view_"] a:active{{
  background:{GREEN} !important;color:#fff !important;outline:2px solid {GREEN_DARK};
  outline-offset:2px;box-shadow:none !important}}
[class*="st-key-view_"] a:hover *,
[class*="st-key-view_"] a:focus *,
[class*="st-key-view_"] a:focus-visible *,
[class*="st-key-view_"] a:active *{{color:#fff !important}}

/* expander "ファイル名をクリック...": màu cố định, không đổi khi hover / click / focus / mở */
[data-testid="stExpander"] details{{background:{card} !important;border:1px solid {border} !important;border-radius:12px}}
[data-testid="stExpander"] summary,
[data-testid="stExpander"] summary:hover,
[data-testid="stExpander"] summary:focus,
[data-testid="stExpander"] summary:focus-visible,
[data-testid="stExpander"] summary:active,
[data-testid="stExpander"] details[open] > summary{{
  background:{card} !important;color:{fg} !important;outline:none !important;box-shadow:none !important;
  -webkit-tap-highlight-color:transparent;border-radius:12px}}
[data-testid="stExpander"] summary *,
[data-testid="stExpander"] summary:hover *,
[data-testid="stExpander"] summary:active *{{color:{fg} !important;fill:{fg} !important}}

.stProgress > div > div > div > div{{background:{GREEN}}}
[data-baseweb="checkbox"] > div:first-child{{background-color:{GREEN} !important}}
</style>"""


def build(items, dicts):
    """items: [(số, tên file)]. Trả về (bytes xlsx, html sheet). Không gọi st.* (chạy nền)."""
    wb = load_workbook(find_template())
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    keys = header_keys(ws)

    r = FIRST_DATA_ROW
    for n, _ in sorted(items):
        dic = dicts[DIC_BY_NUMBER[n]]
        even = [dic.get(k) for k in keys]                 # dòng chẵn: giá trị từ dic
        odd = [ODD_DEFAULTS.get(k) for k in keys]         # dòng lẻ: mặc định
        for vals in (even, odd):
            for col, v in enumerate(vals, start=1):
                if v is not None:
                    ws.cell(row=r, column=col, value=v)
            r += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), sheet_to_html(ws, r - 1)


def make(items):
    return build(items, load_dicts(find_dic())) if items else (None, None)


def parse_value(t):
    t = (t or '').strip()
    if t == '':
        return None
    if re.fullmatch(r'-?(0|[1-9]\d*)', t):
        return int(t)
    if re.fullmatch(r'-?(0|[1-9]\d*)\.\d+', t):
        return float(t)
    return t


def apply_edits(data, edits):
    """Ghi các ô người dùng đã sửa (chỉ dòng lẻ) vào file xlsx."""
    if not edits:
        return data
    wb = load_workbook(io.BytesIO(data))
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    for key, text in edits.items():
        r, c = (int(x) for x in key.split(','))
        if is_editable_row(r):
            ws.cell(row=r, column=c, value=parse_value(text))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@st.cache_data(show_spinner=False, max_entries=8)
def export(data, edits_json):
    return apply_edits(data, json.loads(edits_json))


def main():
    st.set_page_config(page_title='PDF → KDパンチ入力シート', page_icon='📗', layout='wide')

    c1, c2 = st.columns([6, 1])
    dark = c2.toggle('🌙 ダークモード', key='dark')
    st.markdown(app_css(dark), unsafe_allow_html=True)
    c1.title('📗 PDF → KDパンチ入力シート')

    files = st.file_uploader('PDFファイルを1件または複数選択', type='pdf',
                             accept_multiple_files=True)

    if files:
        with st.expander('👁 ファイル名をクリックして内容を確認', expanded=False):
            cols = st.columns(min(len(files), 4))
            for i, f in enumerate(files):
                col = cols[i % len(cols)]
                col.link_button(f'📄 {f.name}', pdf_media_url(f, col),
                                key=f'view_{i}', use_container_width=True)

    if st.button('処理を実行', type='primary', disabled=not files):
        items, errors = [], []
        for f in files:
            n = file_number(f.name)
            if n is None:
                errors.append(f'{f.name}: ファイル名に番号が見つかりません')
            elif n not in DIC_BY_NUMBER:
                errors.append(f'{f.name}: 番号 {n} に対応する辞書がありません')
            else:
                items.append((n, f.name))

        st.session_state['ver'] = st.session_state.get('ver', 0) + 1   # kết quả mới -> xoá chỉnh sửa cũ
        st.session_state['edits'] = {}

        total = SECONDS_PER_FILE * len(files)
        bar = st.progress(0.0, text='処理中...')
        try:
            with ThreadPoolExecutor(max_workers=1) as ex, st.spinner('データを読み込んでいます...'):
                fut = ex.submit(make, items)              # tính toán chạy nền song song với thanh chờ
                for s in range(1, total + 1):
                    time.sleep(1)
                    bar.progress(s / total, text=f'処理中... {s}/{total}s')
                data, sheet = fut.result()
            st.session_state['result'] = (data, sheet, errors)
        except Exception as e:
            st.session_state['result'] = None
            st.error(f'エラー: {e!r}')
        bar.empty()

    res = st.session_state.get('result')
    if res:
        data, sheet, errors = res
        for e in errors:
            st.warning(e)
        if data is not None:
            ver = st.session_state.get('ver', 0)
            head_l, head_r = st.columns([5, 2], vertical_alignment='center')
            head_l.subheader('結果シート')
            edits = kd_sheet(css=excel_css(dark), html=sheet, sig=str(ver), fr=FREEZE_ROWS,
                             init=st.session_state.get('edits', {}), key=f'sheet_{ver}', default={}) or {}
            st.session_state['edits'] = edits
            # nút tải nằm cùng dòng với tiêu đề; file xuất ra đã gồm các ô người dùng sửa
            head_r.download_button('⬇ XLSXをダウンロード', export(data, json.dumps(edits, sort_keys=True)),
                                   file_name=RESULT_NAME, use_container_width=True,
                                   mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    components.html(TOP_BTN, height=0)


if __name__ == '__main__':
    main()


# python -m streamlit run app.py