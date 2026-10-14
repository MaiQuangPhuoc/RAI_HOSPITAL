from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook

path = Path(r'D:\rikai\RAI-LLM\data\sheet_kq.xlsx')

wb = load_workbook(path, data_only=True)
ws = wb['入力シート']


def fmt(v):
    """datetime -> 'tháng/ngày' (vd 2026-04-02 -> '4/2'), giá trị khác giữ nguyên."""
    if isinstance(v, (datetime, date)):
        return f'{v.month}/{v.day}'
    return v


def show(lst):
    """In list với chuỗi hiển thị nguyên ký tự (không bị \\u3000)."""
    items = [f"'{x}'" if isinstance(x, str) else str(x) for x in lst]
    return '[' + ', '.join(items) + ']'


# ---- 1. Dòng chẵn 4,6,8,10,12 (đã đổi ngày) ----
values = {r: [fmt(c.value) for c in ws[r]] for r in range(4, 14, 2)}
value_4, value_6, value_8, value_10, value_12 = (values[r] for r in (4, 6, 8, 10, 12))

# ---- 2. Tên cột ----
row2 = [c.value for c in ws[2]]
row3 = [c.value for c in ws[3]]
zip_idx = next(i for i, v in enumerate(row2) if v == '郵便番号')

name_col = []
for i in range(len(row2)):
    if i <= zip_idx:
        name_col.append(row2[i] if row2[i] is not None else row3[i])
    else:
        name_col.append(row2[i])

# # ---- In có nhãn ----
# print("name_col = ", show(name_col))
# print('=' * 60)


# for r in (4, 6, 8, 10, 12):
#     v = values[r]
#     print('=' * 60)
#     print(f'value_{r}  (giá trị dòng {r}, gồm {len(v)} phần tử)')
#     print(' ', show(v))


# ---- 3. Tạo key duy nhất (tên trùng thì thêm _2, _3...) ----
def make_unique(names):
    count, out = {}, []
    for n in names:
        count[n] = count.get(n, 0) + 1
        out.append(n if count[n] == 1 else f'{n}_{count[n]}')
    return out

keys = make_unique(name_col)

# ---- 4. Tạo dict {tên cột: giá trị} cho từng dòng ----
ROWS = (4, 6, 8, 10, 12)
dicts = {r: dict(zip(keys, values[r])) for r in ROWS}
dict_4, dict_6, dict_8, dict_10, dict_12 = (dicts[r] for r in ROWS)

# ---- 5. In có nhãn ----
def one_line(x):
    return str(x).replace('\n', ' ')      # chỉ để hiển thị, không đổi dữ liệu

for r in ROWS:
    d = dicts[r]
    print('=' * 70)
    print(f'dict_{r}  (dòng {r}, gồm {len(d)} cặp  tên cột: giá trị)')
    for k, v in d.items():
        print(f'  {one_line(k):<28}: {one_line(v)}')


import json

def py_literal(x):
    """Ghi giá trị ra dạng Python: None giữ là None, chuỗi để nguyên chữ Nhật,
    xuống dòng trong chuỗi được ghi thành \\n để mỗi cặp vẫn nằm trên 1 dòng."""
    if x is None:
        return 'None'
    if isinstance(x, str):
        return json.dumps(x, ensure_ascii=False)
    return str(x)          # số, hoặc '4/2' đã là str nên đi vào nhánh trên

out_path = Path(__file__).parent / 'dic_output.txt'

with open(out_path, 'w', encoding='utf-8') as f:
    for r in ROWS:
        f.write(f'dic_{r} = {{\n')
        for k, v in dicts[r].items():
            f.write(f'    {py_literal(k)}: {py_literal(v)},\n')
        f.write('}\n\n')

print('Đã lưu:', out_path)